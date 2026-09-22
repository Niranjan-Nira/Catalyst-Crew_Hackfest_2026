"""
Document parser module for the RAG Knowledge Base.
Supports extraction of clean, structured text and markdown from:
  - PDF documents (.pdf)
  - Microsoft Word documents (.docx, .doc)
  - Plain text & Markdown (.txt, .md)
  - Structured data (.csv, .json)
"""

import io
import json
import re
import csv
from pathlib import Path
from typing import Union, Optional, List


def extract_text_from_pdf(file_input: Union[bytes, io.BytesIO, Path, str]) -> str:
    """
    Extracts text from a PDF document, preserving page demarcations
    as markdown headers (## Page X) to maintain chunking context.
    Uses pypdf with fallback to pypdfium2.
    """
    extracted_pages: List[str] = []

    # Strategy 1: pypdf
    try:
        from pypdf import PdfReader

        if isinstance(file_input, (str, Path)):
            reader = PdfReader(str(file_input))
        elif isinstance(file_input, bytes):
            reader = PdfReader(io.BytesIO(file_input))
        else:
            # Assumed to be BytesIO or file-like object
            reader = PdfReader(file_input)

        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            text = text.strip()
            if text:
                extracted_pages.append(f"## Page {idx + 1}\n\n{text}")

        if extracted_pages:
            return "\n\n".join(extracted_pages)
    except Exception:
        pass

    # Strategy 2: pypdfium2 fallback
    try:
        import pypdfium2 as pdfium

        if isinstance(file_input, (str, Path)):
            doc = pdfium.PdfDocument(str(file_input))
        elif isinstance(file_input, bytes):
            doc = pdfium.PdfDocument(file_input)
        elif hasattr(file_input, "getvalue"):
            doc = pdfium.PdfDocument(file_input.getvalue())
        elif hasattr(file_input, "read"):
            doc = pdfium.PdfDocument(file_input.read())
        else:
            doc = pdfium.PdfDocument(file_input)

        extracted_pages = []
        for idx, page in enumerate(doc):
            textpage = page.get_textpage()
            text = textpage.get_text_range() or ""
            text = text.strip()
            if text:
                extracted_pages.append(f"## Page {idx + 1}\n\n{text}")

        if extracted_pages:
            return "\n\n".join(extracted_pages)
    except Exception:
        pass

    return "\n\n".join(extracted_pages)


def extract_text_from_docx(file_input: Union[bytes, io.BytesIO, Path, str]) -> str:
    """
    Extracts text, headings, bullet lists, and tables from a .docx file,
    formatting headings as markdown (#, ##, ###) and tables as markdown tables.
    """
    import docx

    if isinstance(file_input, (str, Path)):
        doc = docx.Document(str(file_input))
    elif isinstance(file_input, bytes):
        doc = docx.Document(io.BytesIO(file_input))
    else:
        doc = docx.Document(file_input)

    content_blocks: List[str] = []

    # Process document body paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        style_name = p.style.name.lower() if p.style and p.style.name else ""
        if "title" in style_name:
            content_blocks.append(f"# {text}")
        elif "heading 1" in style_name:
            content_blocks.append(f"# {text}")
        elif "heading 2" in style_name:
            content_blocks.append(f"## {text}")
        elif "heading 3" in style_name:
            content_blocks.append(f"### {text}")
        elif "heading 4" in style_name:
            content_blocks.append(f"#### {text}")
        elif "bullet" in style_name or "list" in style_name:
            content_blocks.append(f"- {text}")
        else:
            content_blocks.append(text)

    # Process tables in document
    for table_idx, table in enumerate(doc.tables):
        table_lines: List[str] = []
        rows = list(table.rows)
        if not rows:
            continue

        header_cells = [cell.text.replace("\n", " ").strip() for cell in rows[0].cells]
        if any(header_cells):
            table_lines.append("| " + " | ".join(header_cells) + " |")
            table_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

            for row in rows[1:]:
                row_cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
                table_lines.append("| " + " | ".join(row_cells) + " |")

            if table_lines:
                content_blocks.append("\n" + "\n".join(table_lines) + "\n")

    return "\n\n".join(content_blocks)


def extract_text_from_doc(file_input: Union[bytes, io.BytesIO, Path, str], filename: Optional[str] = None) -> str:
    """
    Extracts text from legacy .doc files.
    First tries python-docx (many .doc files are modern OOXML).
    Then tries win32com if Word is installed.
    Finally falls back to string extraction from raw bytes.
    """
    # Attempt 1: python-docx
    try:
        return extract_text_from_docx(file_input)
    except Exception:
        pass

    # Attempt 2: win32com automation on Windows if Word is installed
    if isinstance(file_input, (str, Path)):
        abs_path = Path(file_input).resolve()
        if abs_path.exists():
            try:
                import win32com.client
                word = win32com.client.Dispatch("Word.Application")
                word.Visible = False
                doc = word.Documents.Open(str(abs_path))
                text = doc.Content.Text
                doc.Close(False)
                word.Quit()
                if text and text.strip():
                    return text.strip()
            except Exception:
                pass

    # Attempt 3: Printable byte strings extraction fallback
    try:
        if isinstance(file_input, (str, Path)):
            raw_data = Path(file_input).read_bytes()
        elif isinstance(file_input, bytes):
            raw_data = file_input
        elif hasattr(file_input, "getvalue"):
            raw_data = file_input.getvalue()
        elif hasattr(file_input, "read"):
            raw_data = file_input.read()
        else:
            raw_data = b""

        # Extract ASCII sequences of len >= 4
        ascii_strings = re.findall(rb"[\x20-\x7E]{4,}", raw_data)
        extracted = [s.decode("ascii", errors="ignore").strip() for s in ascii_strings if len(s.strip()) > 3]
        return "\n".join(extracted)
    except Exception:
        return ""


def extract_text_from_file_data(
    data: Union[str, bytes, io.BytesIO],
    filename: str
) -> str:
    """
    Universal dispatcher that converts any supported document type
    (PDF, Word DOCX/DOC, Markdown, Text, CSV, JSON) into clean, indexable text.
    Leaves all existing text formats working as is.
    """
    ext = Path(filename).suffix.lower()

    # 1. PDF
    if ext == ".pdf":
        if isinstance(data, str):
            data = data.encode("utf-8", errors="ignore")
        return extract_text_from_pdf(data)

    # 2. Word (.docx)
    if ext == ".docx":
        if isinstance(data, str):
            data = data.encode("utf-8", errors="ignore")
        return extract_text_from_docx(data)

    # 3. Word (.doc)
    if ext == ".doc":
        if isinstance(data, str):
            data = data.encode("utf-8", errors="ignore")
        return extract_text_from_doc(data, filename=filename)

    # 4. Standard text/markdown/CSV/JSON
    if isinstance(data, (bytes, bytearray)):
        try:
            raw_text = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                raw_text = data.decode("utf-16-le")
            except Exception:
                raw_text = data.decode("latin1", errors="ignore")
    elif isinstance(data, io.BytesIO):
        byte_val = data.getvalue()
        try:
            raw_text = byte_val.decode("utf-8")
        except UnicodeDecodeError:
            try:
                raw_text = byte_val.decode("utf-16-le")
            except Exception:
                raw_text = byte_val.decode("latin1", errors="ignore")
    else:
        raw_text = str(data)

    # Format JSON nicely if valid JSON
    if ext == ".json":
        try:
            parsed = json.loads(raw_text)
            return json.dumps(parsed, indent=2)
        except Exception:
            return raw_text

    return raw_text


def extract_text_from_file(file_path: Path) -> str:
    """
    Extracts text from a local file Path based on its extension.
    """
    if not file_path.exists():
        return ""

    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext == ".doc":
        return extract_text_from_doc(file_path)
    else:
        try:
            return file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                return file_path.read_text(encoding="utf-16-le")
            except Exception:
                return file_path.read_text(encoding="latin1", errors="ignore")
